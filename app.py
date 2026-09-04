
import io
import time

import streamlit as st
from openpyxl import load_workbook

from instagram_service import (
    get_instagram_data,
    get_korea_time
)

from excel_service import (
    find_url_column,
    get_or_create_column
)


st.set_page_config(
    page_title="Instagram Engagement Checker",
    page_icon="📊"
)


st.title("Instagram Engagement Checker")

st.write(
    "Excel 파일을 업로드하면 Instagram 게시물의 "
    "좋아요 수와 댓글 수를 자동으로 조회합니다."
)


uploaded_file = st.file_uploader(
    "Excel 파일을 업로드하세요.",
    type=["xlsx"]
)


if uploaded_file is not None:

    st.success(
        f"업로드 완료: {uploaded_file.name}"
    )

    if st.button("데이터 조회 시작"):

        workbook = load_workbook(
            uploaded_file
        )

        sheet = workbook.active


        # URL 열 찾기
        url_column = find_url_column(
            sheet
        )

        if url_column is None:

            st.error(
                "URL 열을 찾을 수 없습니다."
            )

            st.stop()


        # 결과 열 준비
        likes_column = get_or_create_column(
            sheet,
            "좋아요"
        )

        comments_column = get_or_create_column(
            sheet,
            "댓글"
        )

        status_column = get_or_create_column(
            sheet,
            "상태"
        )

        checked_at_column = get_or_create_column(
            sheet,
            "조회시각"
        )


        total_rows = sheet.max_row - 1

        success_count = 0
        fail_count = 0


        progress_bar = st.progress(0)

        status_text = st.empty()


        for row in range(
            2,
            sheet.max_row + 1
        ):

            url = sheet.cell(
                row=row,
                column=url_column
            ).value


            current = row - 1

            status_text.text(
                f"{current} / {total_rows} 처리 중..."
            )


            if not url:

                sheet.cell(
                    row=row,
                    column=status_column,
                    value="EMPTY_URL"
                )

                sheet.cell(
                    row=row,
                    column=checked_at_column,
                    value=get_korea_time()
                )

                fail_count += 1

            else:

                likes, comments, status = (
                    get_instagram_data(url)
                )

                checked_at = get_korea_time()


                sheet.cell(
                    row=row,
                    column=likes_column,
                    value=likes
                )

                sheet.cell(
                    row=row,
                    column=comments_column,
                    value=comments
                )

                sheet.cell(
                    row=row,
                    column=status_column,
                    value=status
                )

                sheet.cell(
                    row=row,
                    column=checked_at_column,
                    value=checked_at
                )


                if status == "SUCCESS":
                    success_count += 1
                else:
                    fail_count += 1


                time.sleep(2)


            progress_bar.progress(
                current / total_rows
            )


        status_text.text(
            "처리 완료"
        )


        st.success(
            "Instagram 데이터 수집이 완료되었습니다."
        )


        col1, col2, col3 = st.columns(3)

        col1.metric(
            "총 URL",
            total_rows
        )

        col2.metric(
            "성공",
            success_count
        )

        col3.metric(
            "실패",
            fail_count
        )


        # 결과 Excel을 메모리에 저장
        output = io.BytesIO()

        workbook.save(
            output
        )

        output.seek(0)


        result_name = (
            uploaded_file.name
            .replace(
                ".xlsx",
                "_result.xlsx"
            )
        )


        st.download_button(
            label="결과 Excel 다운로드",
            data=output,
            file_name=result_name,
            mime=(
                "application/"
                "vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )
        )
