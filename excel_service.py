
from openpyxl import load_workbook


URL_CANDIDATES = [
    "URL",
    "Instagram URL",
    "인스타그램 URL",
    "게시물 URL",
    "링크"
]


def load_excel(file_path):
    workbook = load_workbook(file_path)
    sheet = workbook.active

    return workbook, sheet


def find_url_column(sheet):
    headers = [
        cell.value
        for cell in sheet[1]
    ]

    for i, header in enumerate(headers, start=1):
        if header in URL_CANDIDATES:
            return i

    return None


def get_or_create_column(sheet, header_name):
    headers = [
        cell.value
        for cell in sheet[1]
    ]

    if header_name in headers:
        return headers.index(header_name) + 1

    new_column = sheet.max_column + 1

    sheet.cell(
        row=1,
        column=new_column,
        value=header_name
    )

    return new_column
